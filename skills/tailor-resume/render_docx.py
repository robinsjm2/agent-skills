"""Render a resume written in the tailor-resume Markdown format to .docx.

Usage:
    uv run --with python-docx render_docx.py resume.md [output.docx]

Format (one element per line):
    # Name
    contact line (first plain line after the name)
    ## Section heading
    ### Company, Location | Dates      (dates are right-aligned)
    *Job title*
    - bullet (supports **bold** spans)
    plain paragraph (supports **bold** spans)
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

FONT = "Calibri"
BODY_PT = 10.5
NAME_PT = 18
HEADING_PT = 11
RIGHT_TAB = Inches(7.5)


def add_runs(paragraph, text, size=BODY_PT, bold=False, italic=False):
    """Add text to a paragraph, honoring **bold** spans."""
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not part:
            continue
        run = paragraph.add_run(part)
        run.font.name = FONT
        run.font.size = Pt(size)
        run.bold = bold or i % 2 == 1
        run.italic = italic


def spacing(paragraph, before=0, after=2):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)


def bottom_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, val in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "808080")):
        bottom.set(qn(f"w:{key}"), val)
    borders.append(bottom)
    p_pr.append(borders)


def render(md_path, out_path):
    doc = Document()
    section = doc.sections[0]
    section.left_margin = section.right_margin = Inches(0.5)
    section.top_margin = section.bottom_margin = Inches(0.625)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY_PT)

    after_name = False
    for raw in Path(md_path).read_text().splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("# "):
            p = doc.add_paragraph()
            add_runs(p, line[2:], size=NAME_PT, bold=True)
            spacing(p, after=0)
            after_name = True
        elif after_name:
            p = doc.add_paragraph()
            add_runs(p, line)
            spacing(p, after=4)
            after_name = False
        elif line.startswith("## "):
            p = doc.add_paragraph()
            add_runs(p, line[3:].upper(), size=HEADING_PT, bold=True)
            spacing(p, before=8, after=3)
            bottom_border(p)
        elif line.startswith("### "):
            left, _, dates = line[4:].partition(" | ")
            p = doc.add_paragraph()
            p.paragraph_format.tab_stops.add_tab_stop(RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT)
            add_runs(p, left, bold=True)
            if dates:
                add_runs(p, "\t" + dates, bold=True)
            spacing(p, before=6, after=0)
        elif line.startswith("*") and line.endswith("*") and not line.startswith("**"):
            p = doc.add_paragraph()
            add_runs(p, line.strip("*"), italic=True)
            spacing(p, after=2)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, line[2:])
            spacing(p, after=1)
        else:
            p = doc.add_paragraph()
            add_runs(p, line)
            spacing(p, after=2)

    doc.save(out_path)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    src = Path(sys.argv[1])
    dest = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_suffix(".docx")
    render(src, dest)
    print(dest)
